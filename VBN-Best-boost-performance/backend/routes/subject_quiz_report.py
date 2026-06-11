from extension import db,mail
from flask_mail import Message
from openpyxl import Workbook
from models.user_model import User
from models.subject_model import Subject
from models.test_result_model import TestResult
from models.school_model import School
from datetime import datetime
import os
import json


def _parse_subject_ids(selected_subjects):

    if not selected_subjects:
        return []

    text = str(selected_subjects).strip()

    try:
        parsed = json.loads(text)

        if isinstance(parsed, list):
            return [int(x) for x in parsed if str(x).strip()]

    except Exception:
        pass

    ids = []

    for x in text.split(","):

        token = x.strip()

        if token.isdigit():
            ids.append(int(token))

    return ids

def send_subject_quiz_report(app):

    with app.app_context():

        today = datetime.utcnow().date()

        teachers=User.query.filter_by(
            position=3
        ).all()
        print("Teachers found:", len(teachers))

        for teacher in teachers:

            if not teacher.email:
                continue

            ids = _parse_subject_ids(
                teacher.selected_subjects
            )

            if not ids:
                print(
                    "No valid selected subjects for teacher:",
                    teacher.id
                )
                continue

            for subject_id in ids:
                subject = Subject.query.get(
                subject_id)

                if not subject:
                    continue

                results=(
                    TestResult.query
                    .filter(
                        TestResult.subject_id
                        ==subject_id,

                        TestResult.test_type
                        =="subject",

                        db.func.date(
                            TestResult.taken_at
                        )==today
                    )
                    .all()
                )
                print(
                "Subject:",
                subject_id,
                "Results:",
                len(results)
            )
                

                wb=Workbook()

                ws1=wb.active
                ws1.title="Attended"

                ws1.append([
                    "Student",
                    "Score %",
                    "Correct",
                    "Total"
                ])
                if not results:
                    ws1.append([
                        "No students attended today's test"
                    ])

                attended=[]

                for r in results:

                    user=User.query.get(
                        r.user_id
                    )

                    if not user:
                        continue

                    attended.append(
                        user.id
                    )

                    ws1.append([

                        user.name,
                        r.score_percent,
                        r.correct_answers,
                        r.total_questions

                    ])

                ws2=wb.create_sheet(
                    "Not Attended"
                )

                ws2.append([
                    "Student"
                ])

                students=User.query.filter_by(
                    school_id=
                    teacher.school_id,

                    position=2
                ).all()

                for s in students:

                    if s.id not in attended:

                        ws2.append([
                            s.name
                        ])

                ws3=wb.create_sheet(
                    "Summary"
                )
                total_students = len(students)

                ws3.append([
                    "Subject",
                    subject.subject_name
                ])

                ws3.append([
                    "Total Students",
                    total_students
                ])

                ws3.append([
                    "Attended",
                    len(attended)
                ])

                ws3.append([
                    "Not Attended",
                    total_students - len(attended)
                ])

                attendance_percent = round(
                    (len(attended) / total_students) * 100,
                    1
                ) if total_students else 0

                ws3.append([
                    "Attendance %",
                    f"{attendance_percent}%"
                ])

                filename=(
                    f"subject_"
                    f"{subject.id}.xlsx"
                )

                wb.save(
                    filename
                )

                school=School.query.get(
                    teacher.school_id
                )

                school_name=(
                    school.name
                    if school else "School"
                )

                msg=Message(
                    subject=
                    f"Subject Test ({school_name})",

                    recipients=[
                        teacher.email
                    ]
                )

                msg.body = f"""
Respected Teacher,

We hope this message finds you well.

Please find attached the subject test report for today.

Subject: {subject.subject_name}

Total Students: {total_students}
Attended: {len(attended)}
Not Attended: {total_students - len(attended)}
Attendance: {attendance_percent}%

Thank you for your continued support and dedication.

Yours sincerely,
VBN Boost Performance Team
"""

                try:
                    with open(
                        filename,
                        "rb"
                    ) as fp:

                        msg.attach(
                            filename,
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            fp.read()
                        )

                    mail.send(msg)

                    print(
                        "Subject report mail sent to:",
                        teacher.email,
                        "subject:",
                        subject_id
                    )

                except Exception as e:
                    print(
                        "Subject report mail failed:",
                        str(e),
                        "teacher:",
                        teacher.id,
                        "subject:",
                        subject_id
                    )

                finally:
                    if os.path.exists(filename):
                        os.remove(filename)